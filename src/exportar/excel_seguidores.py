"""Fase 3: Excel de los seguidores del perfil (foto, nombre, seguidos, seguidores y PRIMER DÍGITO para Benford).

Se usa desde main.py:   python main.py excel <usuario>

Entrada : raw/seguidores_lista.json (Fase 1) + raw/seguidores_contadores.json (Fase 2)
Salida  : raw/instagram_seguidores_stats.xlsx
    Hoja "Seguidores": Foto | Usuario | Nombre | Seguidos | Seguidores | Primer dígito | Estado | Fecha consulta
    Hoja "Benford"   : frecuencia real vs Benford, chi-cuadrada, valor p, MAD y veredicto (todo con fórmulas)
    Hoja "Resumen"   : conteos por estado y cantidad de datos válidos para Benford (con fórmulas)

La columna "Primer dígito" es una FÓRMULA de Excel sobre "Seguidores" (se ve y se audita en la propia celda);
queda vacía si no hay dato. En el archivo las funciones se guardan en inglés (IF, LEFT...); Excel en español
las muestra como SI, IZQUIERDA...
"""
import io
import json
import os
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

from openpyxl.formatting.rule import FormulaRule

from src.benford.benford import (ALFA_BOT, ALFA_HUMANO, GRADOS_LIBERTAD, MIN_ABSOLUTO, MIN_MUESTRA, analizar,
                                 seguidores_validos)
from src.config import rutas_usuario

ZONA = ZoneInfo("America/Guayaquil")  # Ecuador (UTC-5)
MINIATURA_PX = 44          # lado de la foto dentro de la celda
ALTO_FILA_PT = 36          # alto de fila en puntos (~48 px)
MIN_BENFORD = MIN_ABSOLUTO  # por debajo no hay veredicto; entre 30 y 109 el valor p se calcula por simulación

COLUMNAS = ["Foto", "Usuario", "Nombre", "Seguidos", "Seguidores", "Primer dígito", "Estado", "Fecha consulta"]
ANCHOS = [8, 24, 28, 12, 13, 14, 12, 19]
COL_SEGUIDORES = "E"
COL_DIGITO = "F"
COL_ESTADO = "G"

AZUL = "1F3A5F"
ESTILO_ENCABEZADO = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=AZUL),
                         alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
BORDE = Border(bottom=Side(style="thin", color="D9D9D9"))


def _leer_json(ruta):
    if not os.path.exists(ruta):
        return None
    try:
        with open(ruta, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return None


def _miniatura(ruta_foto):
    """Foto -> PNG cuadrado pequeño en memoria (liviano para el Excel). None si no se puede leer."""
    try:
        with Image.open(ruta_foto) as im:
            im = im.convert("RGBA")
            lado = min(im.size)                     # recorte central cuadrado
            izq, arr = (im.width - lado) // 2, (im.height - lado) // 2
            im = im.crop((izq, arr, izq + lado, arr + lado)).resize((MINIATURA_PX, MINIATURA_PX), Image.LANCZOS)
            buf = io.BytesIO()
            im.save(buf, format="PNG", optimize=True)
            buf.seek(0)
            return buf
    except Exception:
        return None


def unir_datos(lista_json, contadores_json):
    """Una fila por seguidor de la lista (mismo orden), con sus contadores si ya se consultaron."""
    cont = {r["usuario"]: r for r in (contadores_json or {}).get("resultados", [])}
    filas, vistos = [], set()
    for s in lista_json.get("seguidores", []):
        u = s["usuario"]
        if u in vistos:
            continue
        vistos.add(u)
        c = cont.get(u, {})
        filas.append({
            "usuario": u,
            "nombre": s.get("nombre") or "",
            "url": s.get("url") or f"https://www.instagram.com/{u}/",
            "foto_archivo": s.get("foto_archivo"),
            "seguidos": c.get("seguidos"),
            "seguidores": c.get("seguidores"),
            "estado": c.get("estado", "pendiente"),
            "fecha_consulta": c.get("fecha_consulta") or "",
        })
    return filas


def _formula_digito(fila):
    # IF(el valor es un número > 0, primer carácter convertido a número, vacío)
    return f'=IF(AND(ISNUMBER({COL_SEGUIDORES}{fila}),{COL_SEGUIDORES}{fila}>0),VALUE(LEFT({COL_SEGUIDORES}{fila},1)),"")'


def _hoja_benford(wb, ultima, p_simulado=None):
    """Hoja con la prueba completa en fórmulas: se recalcula sola si cambias datos en la hoja Seguidores."""
    bf = wb.create_sheet("Benford", 1)
    rango_d = f"Seguidores!${COL_DIGITO}$2:${COL_DIGITO}${max(ultima, 2)}"
    bf.append(["Dígito", "Frecuencia real", "% Real", "% Benford", "Esperado", "(O-E)²/E", "|Real-Benford|"])
    for d in range(1, 10):
        f = d + 1
        bf.append([d, f"=COUNTIF({rango_d},A{f})", f"=IF($B$13>0,B{f}/$B$13,0)", f"=LOG10(1+1/A{f})",
                   f"=$B$13*D{f}", f"=IF(E{f}>0,(B{f}-E{f})^2/E{f},0)", f"=ABS(C{f}-D{f})"])
    bf.append(["Total", "=SUM(B2:B10)", "=SUM(C2:C10)", "=SUM(D2:D10)", "=SUM(E2:E10)", "=SUM(F2:F10)", None])
    bf.append([])
    resumen = [
        ("Datos válidos (n)", "=B11"),
        ("Chi-cuadrada", "=F11"),
        ("Grados de libertad", GRADOS_LIBERTAD),
        ("Valor p", "=CHIDIST(B14,B15)"),
        ("Crítico al 5 %", "=CHIINV(0.05,B15)"),
        ("Crítico al 1 %", "=CHIINV(0.01,B15)"),
        ("MAD (Nigrini)", "=SUM(G2:G10)/9"),
        ("Nivel MAD", '=IF(B19<=0.006,"conformidad cercana",IF(B19<=0.012,"conformidad aceptable",'
                      'IF(B19<=0.015,"conformidad marginal","no conforme")))'),
        ("Mínimo de datos para veredicto", MIN_BENFORD),
        ("VEREDICTO", f'=IF(B13<B21,"MUESTRA INSUFICIENTE",IF(IF(B13>={MIN_MUESTRA},B16,B24)>={ALFA_HUMANO},"HUMANO",'
                      f'IF(IF(B13>={MIN_MUESTRA},B16,B24)>={ALFA_BOT},"POSIBLE BOT","BOT")))'),
        ("Método del valor p", f'=IF(B13>={MIN_MUESTRA},"chi-cuadrada (fórmula, en vivo)",'
                               f'"simulación Monte Carlo (valor fijado al generar el Excel)")'),
        ("Valor p por simulación (solo n < %d)" % MIN_MUESTRA, p_simulado if p_simulado is not None else ""),
    ]
    for etiqueta, valor in resumen:
        bf.append([etiqueta, valor])
    bf.append([])
    bf.append(["Reglas: p ≥ 0.05 → HUMANO · 0.01 ≤ p < 0.05 → POSIBLE BOT · p < 0.01 → BOT. "
               "Con menos de %d datos el valor p se toma de la simulación (fila 24, fija: si cambias datos aquí, vuelve a ejecutar "
               "'python main.py excel'). El MAD es solo referencial con menos de 1000 datos (decide la chi-cuadrada). " % MIN_MUESTRA +
               "Los seguidores con 0 seguidores no tienen primer dígito y no cuentan. "
               "El veredicto describe a la audiencia en conjunto; es un indicio estadístico, no una prueba."])

    for celda in bf[1]:
        celda.font, celda.fill, celda.alignment = (ESTILO_ENCABEZADO["font"], ESTILO_ENCABEZADO["fill"],
                                                   ESTILO_ENCABEZADO["alignment"])
    for ancho, col in zip((30, 16, 11, 11, 11, 11, 15), "ABCDEFG"):
        bf.column_dimensions[col].width = ancho
    for f in range(2, 12):
        for col, formato in (("C", "0.00%"), ("D", "0.00%"), ("E", "0.0"), ("F", "0.000"), ("G", "0.0000")):
            bf[f"{col}{f}"].number_format = formato
    for f in (14,):
        bf[f"B{f}"].number_format = "0.000"
    bf["B16"].number_format = "0.0000"
    bf["B17"].number_format = bf["B18"].number_format = "0.000"
    bf["B19"].number_format = "0.0000"
    for f in range(2, 11):
        bf[f"A{f}"].alignment = Alignment(horizontal="center")
    for celda in bf[11]:
        celda.font = Font(bold=True)
    bf["B24"].number_format = "0.0000"
    for f in range(13, 25):
        bf[f"A{f}"].font = Font(bold=True)
        bf[f"B{f}"].alignment = Alignment(horizontal="left")
    bf["A26"].alignment = Alignment(wrap_text=True, vertical="top")
    bf.merge_cells("A26:G29")

    # Color del veredicto
    colores = (("HUMANO", "C6EFCE", "006100"), ("POSIBLE BOT", "FFEB9C", "9C5700"),
               ("BOT", "FFC7CE", "9C0006"), ("MUESTRA INSUFICIENTE", "E7E6E6", "404040"))
    for texto, fondo, letra in colores:
        bf.conditional_formatting.add("B22", FormulaRule(
            formula=[f'$B$22="{texto}"'], fill=PatternFill("solid", bgColor=fondo, fgColor=fondo),
            font=Font(bold=True, color=letra)))
    bf["B22"].font = Font(bold=True)
    return bf


def construir_excel(usuario, filas, carpeta_raw, ruta_salida, p_simulado=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Seguidores"
    ws.append(COLUMNAS)
    for i, ancho in enumerate(ANCHOS, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for celda in ws[1]:
        celda.font, celda.fill, celda.alignment = (ESTILO_ENCABEZADO["font"], ESTILO_ENCABEZADO["fill"],
                                                   ESTILO_ENCABEZADO["alignment"])
    ws.row_dimensions[1].height = 28

    buffers, con_foto = [], 0                       # los buffers deben vivir hasta guardar el libro
    for n, f in enumerate(filas, start=2):
        ws.append([None, f["usuario"], f["nombre"], f["seguidos"], f["seguidores"], _formula_digito(n),
                   f["estado"], f["fecha_consulta"]])
        ws.row_dimensions[n].height = ALTO_FILA_PT

        c = ws[f"B{n}"]                              # usuario con enlace al perfil
        c.hyperlink, c.font = f["url"], Font(color="0563C1", underline="single")

        for col in "DE":
            ws[f"{col}{n}"].number_format = "#,##0"
        for col in "ABCDEFGH":
            ws[f"{col}{n}"].border = BORDE
            ws[f"{col}{n}"].alignment = Alignment(vertical="center",
                                                  horizontal="center" if col in "AFG" else None)

        if f["foto_archivo"]:
            buf = _miniatura(os.path.join(carpeta_raw, f["foto_archivo"]))
            if buf is not None:
                buffers.append(buf)
                img = XLImage(buf)
                img.width = img.height = MINIATURA_PX
                ws.add_image(img, f"A{n}")
                con_foto += 1

    ultima = len(filas) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{max(ultima, 2)}"

    # ---- Hoja Resumen (todo con fórmulas: se actualiza si editas la hoja Seguidores) ----
    rs = wb.create_sheet("Resumen")
    rango_u = f"Seguidores!$B$2:$B${max(ultima, 2)}"
    rango_e = f"Seguidores!${COL_ESTADO}$2:${COL_ESTADO}${max(ultima, 2)}"
    rango_d = f"Seguidores!${COL_DIGITO}$2:${COL_DIGITO}${max(ultima, 2)}"
    datos = [
        ("Perfil analizado", f"@{usuario}"),
        ("Fecha de generación", datetime.now(ZONA).strftime("%Y-%m-%d %H:%M")),
        ("Seguidores en la lista", f"=COUNTA({rango_u})"),
        ("Con contadores (ok)", f'=COUNTIF({rango_e},"ok")'),
        ("Cuenta privada (contadores visibles)", f'=COUNTIF({rango_e},"privado")'),
        ("Perfil inexistente", f'=COUNTIF({rango_e},"no_existe")'),
        ("Con error (se reintentan)", f'=COUNTIF({rango_e},"error")'),
        ("Pendientes de consultar", f'=COUNTIF({rango_e},"pendiente")'),
        ("Datos válidos para Benford (primer dígito 1-9)", f"=COUNT({rango_d})"),
        ("¿Muestra suficiente?", f'=IF(B10>={MIN_MUESTRA},"Sí (chi-cuadrada clásica)",IF(B10>={MIN_BENFORD},'
                                 f'"Reducida: valor p por simulación","No: mínimo {MIN_BENFORD} para veredicto")'),
        ("Cuentas con 0 seguidores (sin primer dígito)", f'=COUNTIF(Seguidores!${COL_SEGUIDORES}$2:${COL_SEGUIDORES}${max(ultima, 2)},0)'),
        ("Fotos incrustadas", con_foto),
    ]
    rs.append(["Concepto", "Valor"])
    for fila in datos:
        rs.append(list(fila))
    for celda in rs[1]:
        celda.font, celda.fill, celda.alignment = (ESTILO_ENCABEZADO["font"], ESTILO_ENCABEZADO["fill"],
                                                   ESTILO_ENCABEZADO["alignment"])
    rs.column_dimensions["A"].width = 48
    rs.column_dimensions["B"].width = 34
    for fila in rs.iter_rows(min_row=2, max_col=2):
        fila[1].alignment = Alignment(horizontal="left")

    _hoja_benford(wb, ultima, p_simulado)

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    wb.save(ruta_salida)
    return {"filas": len(filas), "con_foto": con_foto}


def generar_excel_seguidores(usuario):
    rutas = rutas_usuario(usuario)
    lista = _leer_json(rutas.seguidores_json)
    if not lista or not lista.get("seguidores"):
        raise SystemExit(f"No hay lista de seguidores en {rutas.raw}. Ejecuta primero:  "
                         f"python main.py seguidores {usuario}")
    contadores = _leer_json(rutas.contadores_json)
    if not contadores:
        print("[Excel] AVISO: aún no hay contadores (Fase 2); las columnas Seguidos/Seguidores quedarán vacías. "
              f"Ejecuta:  python main.py contadores {usuario}")

    filas = unir_datos(lista, contadores)
    p_sim = analizar(seguidores_validos(contadores))["valor_p_simulado"] if contadores else None
    try:
        res = construir_excel(usuario, filas, rutas.raw, rutas.seguidores_xlsx, p_simulado=p_sim)
    except PermissionError:
        raise SystemExit(f"No se pudo escribir {rutas.seguidores_xlsx}: ciérralo en Excel y vuelve a intentar.")

    validos = sum(1 for f in filas if isinstance(f["seguidores"], int) and f["seguidores"] > 0)
    print("\n--- EXCEL ---")
    print(f"filas: {res['filas']}  |  con foto: {res['con_foto']}  |  con seguidores válidos para Benford: {validos}")
    if validos < MIN_BENFORD:
        print(f"AVISO: hay {validos} datos válidos; para emitir veredicto se necesitan al menos {MIN_BENFORD}.")
    elif validos < MIN_MUESTRA:
        print(f"Nota: {validos} datos válidos (< {MIN_MUESTRA}): el valor p del Excel se toma de la simulación Monte Carlo.")
    print(f"archivo: {rutas.seguidores_xlsx}")
    return res
