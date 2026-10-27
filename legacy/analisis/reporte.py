"""Documento Word + gráfico radar del perfil Big Five (se guardan en la carpeta del scraping)."""
import sys
import os

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.dirname(CURRENT_DIR)
for ruta in (ROOT_DIR, CURRENT_DIR):
    if ruta not in sys.path:
        sys.path.insert(0, ruta)

import json
import unicodedata
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from config import DATOS_DIR, INSTAGRAM_TARGET

ORDEN = ["apertura", "responsabilidad", "extraversion", "amabilidad", "neuroticismo"]
NOMBRES = {
    "apertura": "Apertura a la experiencia",
    "responsabilidad": "Responsabilidad",
    "extraversion": "Extraversión",
    "amabilidad": "Amabilidad",
    "neuroticismo": "Neuroticismo",
}
DEFINICIONES = {
    "apertura": "Curiosidad, imaginación e interés por el arte, las ideas nuevas y las experiencias distintas.",
    "responsabilidad": "Organización, constancia, disciplina y orientación a metas y compromisos.",
    "extraversion": "Energía social, sociabilidad y tendencia a buscar la compañía y la estimulación de otros.",
    "amabilidad": "Empatía, cooperación, confianza y preocupación por los demás.",
    "neuroticismo": "Tendencia a experimentar emociones negativas como ansiedad, tristeza o estrés (lo opuesto es la estabilidad emocional).",
}
COLOR = RGBColor(0x1F, 0x3A, 0x5F)


def _clave(nombre):
    """'Extraversión ' -> 'extraversion'"""
    n = unicodedata.normalize("NFD", nombre or "").encode("ascii", "ignore").decode().lower().strip()
    return n


def _nivel(puntaje):
    return "Bajo" if puntaje < 40 else "Medio" if puntaje <= 60 else "Alto"


def _fecha_scraping(carpeta):
    try:
        return datetime.strptime(carpeta, "%Y-%m-%d_%H%M").strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return carpeta


def generar_radar(rasgos, ruta_png, usuario):
    valores = [rasgos[k]["puntaje"] for k in ORDEN]
    etiquetas = [NOMBRES[k] for k in ORDEN]
    ang = np.linspace(0, 2 * np.pi, len(ORDEN), endpoint=False).tolist()
    v, a = valores + valores[:1], ang + ang[:1]

    fig, ax = plt.subplots(figsize=(6, 6), subplot_kw=dict(polar=True))
    ax.set_theta_offset(np.pi / 2)      # Apertura arriba
    ax.set_theta_direction(-1)          # sentido horario
    ax.plot(a, v, color="#1F3A5F", linewidth=2)
    ax.fill(a, v, color="#1F3A5F", alpha=0.25)
    ax.set_xticks(ang)
    ax.set_xticklabels(etiquetas, fontsize=10)
    ax.tick_params(axis='x', pad=14)
    ax.set_ylim(0, 100)
    ax.set_yticks([20, 40, 60, 80, 100])
    ax.set_yticklabels(["20", "40", "60", "80", "100"], fontsize=8, color="gray")
    for ang_i, val in zip(ang, valores):
        ax.text(ang_i, max(val - 9, 6) if val > 85 else val + 8, str(val), ha="center", va="center", fontsize=10, fontweight="bold")
    ax.set_title(f"Perfil Big Five estimado — @{usuario}", fontsize=12, pad=22)
    fig.tight_layout()
    fig.savefig(ruta_png, dpi=200)
    plt.close(fig)


def _titulo(doc, texto, nivel=1):
    h = doc.add_heading(texto, level=nivel)
    for r in h.runs:
        r.font.color.rgb = COLOR
    return h


def _cargar_textos():
    ruta = os.path.join(DATOS_DIR, "metadatos.json")
    if os.path.exists(ruta):
        with open(ruta, encoding="utf-8") as f:
            return {d["id"]: d for d in json.load(f)}
    return {}


def generar_reporte(perfil):
    """perfil: diccionario devuelto por generar_perfil_bigfive(). Devuelve (ruta_docx, ruta_png)."""
    meta = perfil.get("meta", {})
    usuario = meta.get("usuario", INSTAGRAM_TARGET)
    rasgos = {_clave(r["rasgo"]): r for r in perfil["rasgos"]}
    faltan = [k for k in ORDEN if k not in rasgos]
    if faltan:
        raise SystemExit(f"Gemini no devolvió estos rasgos: {faltan}. Vuelve a ejecutar el análisis.")
    textos = _cargar_textos()

    ruta_png = os.path.join(DATOS_DIR, "radar_bigfive.png")
    ruta_docx = os.path.join(DATOS_DIR, "reporte_bigfive.docx")
    generar_radar(rasgos, ruta_png, usuario)

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    t = doc.add_heading("Perfil de personalidad Big Five (OCEAN)", level=0)
    for r in t.runs:
        r.font.color.rgb = COLOR
    p = doc.add_paragraph("Estimación exploratoria a partir de publicaciones de Instagram")
    p.runs[0].italic = True

    tabla = doc.add_table(rows=0, cols=2)
    tabla.style = "Table Grid"
    filas = [
        ("Perfil analizado", f"@{usuario}"),
        ("Scraping realizado", _fecha_scraping(meta.get("scraping", os.path.basename(DATOS_DIR)))),
        ("Textos analizados", str(meta.get("textos_disponibles", "—"))),
        ("Modelo de lenguaje", str(meta.get("modelo", "—"))),
        ("Informe generado", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    for k, v in filas:
        c = tabla.add_row().cells
        c[0].text, c[1].text = k, v
        c[0].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()

    _titulo(doc, "Resultados", 1)
    doc.add_picture(ruta_png, width=Cm(11.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    res = doc.add_table(rows=1, cols=4)
    res.style = "Table Grid"
    for i, h in enumerate(("Rasgo", "Puntaje (0-100)", "Nivel", "Confianza")):
        res.rows[0].cells[i].text = h
        res.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for k in ORDEN:
        r = rasgos[k]
        c = res.add_row().cells
        c[0].text = NOMBRES[k]
        c[1].text = str(r["puntaje"])
        c[2].text = _nivel(r["puntaje"])
        c[3].text = f"{r['confianza']:.2f}"
    nota = doc.add_paragraph("Niveles: bajo (<40), medio (40–60), alto (>60). La confianza va de 0 a 1 y está "
                             f"limitada a {meta.get('techo_confianza', '—')} por la cantidad de texto disponible.")
    nota.runs[0].font.size = Pt(9)
    nota.runs[0].italic = True

    _titulo(doc, "Resumen", 1)
    doc.add_paragraph(perfil.get("resumen", ""))

    _titulo(doc, "Interpretación por rasgo", 1)
    for k in ORDEN:
        r = rasgos[k]
        _titulo(doc, f"{NOMBRES[k]} — {r['puntaje']}/100 ({_nivel(r['puntaje'])})", 2)
        d = doc.add_paragraph(DEFINICIONES[k])
        d.runs[0].italic = True
        doc.add_paragraph(r.get("justificacion", ""))
        citas = []
        for ev in r.get("evidencia", []):
            if ev in textos:
                tx = textos[ev]["texto"].replace("\n", " ")
                tx = tx if len(tx) <= 140 else tx[:137] + "…"
                fecha = f" ({textos[ev]['fecha']})" if textos[ev].get("fecha") else ""
                citas.append(f"{ev}{fecha}: “{tx}”")
            elif ev:
                citas.append(f"Métrica: {ev}")
        if citas:
            doc.add_paragraph("Evidencia citada:").runs[0].bold = True
            for c in citas:
                doc.add_paragraph(c, style="List Bullet")

    _titulo(doc, "Limitaciones y consideraciones éticas", 1)
    doc.add_paragraph(perfil.get("limitaciones", ""))
    for texto in (
        "Es una estimación exploratoria generada por un modelo de lenguaje; no es un diagnóstico ni una medición psicológica validada.",
        "Las publicaciones muestran una faceta pública y puntual de la persona, y abarcan períodos distintos de su vida.",
        "El análisis se realizó con la autorización de la titular de la cuenta; no debe compartirse sin su consentimiento.",
    ):
        doc.add_paragraph(texto, style="List Bullet")

    doc.save(ruta_docx)
    print(f"[Reporte] Documento: {ruta_docx}\n[Reporte] Gráfico:   {ruta_png}")
    return ruta_docx, ruta_png


if __name__ == "__main__":
    ruta = os.path.join(DATOS_DIR, "perfil_bigfive.json")
    if not os.path.exists(ruta):
        raise SystemExit("Aún no hay perfil_bigfive.json; ejecuta antes: python analisis/main.py")
    with open(ruta, encoding="utf-8") as f:
        generar_reporte(json.load(f))
